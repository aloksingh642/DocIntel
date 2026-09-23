"""Text-extractor tests: per-type extraction + graceful corruption handling."""
import pytest
from docx import Document as DocxDocument

from app.extractors import ExtractionError, get_extractor
from app.extractors.pdf_extractor import PdfExtractor


def test_txt_extractor(tmp_path):
    p = tmp_path / "note.txt"
    p.write_text("hello\nworld")
    out = get_extractor(".txt").extract(p)
    assert out.text == "hello\nworld"
    assert out.pages == ["hello\nworld"]


def test_txt_extractor_rejects_empty(tmp_path):
    p = tmp_path / "empty.txt"
    p.write_text("")
    with pytest.raises(ExtractionError):
        get_extractor(".txt").extract(p)


def test_docx_extractor(tmp_path):
    p = tmp_path / "resume.docx"
    doc = DocxDocument()
    doc.add_heading("John Doe", 0)
    doc.add_paragraph("Skills: Python, SQL")
    doc.save(p)
    out = get_extractor(".docx").extract(p)
    assert "John Doe" in out.text
    assert "Python" in out.text


def test_docx_extractor_corrupted(tmp_path):
    p = tmp_path / "broken.docx"
    p.write_bytes(b"PK this is not a real docx at all")
    with pytest.raises(ExtractionError):
        get_extractor(".docx").extract(p)


def test_pdf_extractor_text_layer(tmp_path):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    p = tmp_path / "resume.pdf"
    c = canvas.Canvas(str(p), pagesize=letter)
    c.drawString(100, 750, "John Doe - Python Developer with FastAPI and PostgreSQL experience")
    c.save()

    out = PdfExtractor().extract(p)
    assert "John Doe" in out.text
    assert not out.used_ocr


def test_pdf_extractor_corrupted(tmp_path):
    p = tmp_path / "broken.pdf"
    p.write_bytes(b"%PDF-1.4 but actually garbage bytes here")
    with pytest.raises(ExtractionError):
        PdfExtractor().extract(p)


def test_get_extractor_unknown_extension():
    with pytest.raises(ExtractionError):
        get_extractor(".exe")
