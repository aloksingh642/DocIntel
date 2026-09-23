"""DOCX text extraction (paragraphs + tables)."""
from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError

from app.extractors.base import ExtractionError, ExtractionOutput, TextExtractor


class DocxExtractor(TextExtractor):
    def extract(self, path: Path) -> ExtractionOutput:
        try:
            doc = DocxDocument(str(path))
        except (PackageNotFoundError, Exception) as exc:
            raise ExtractionError(f"Unable to read DOCX: {exc.__class__.__name__}") from exc

        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip(" |"):
                    parts.append(row_text)

        text = "\n".join(parts).strip()
        if not text:
            raise ExtractionError("DOCX contains no extractable text.")
        return ExtractionOutput(text=text, pages=[text], metadata={"format": "docx"})
